"""Pre-export fidelity linter + fixer for layered canvas → PPTX.

Scene units are inches. Text is measured with the font file PowerPoint will actually use on the viewer's
machine (the brand font if installed, otherwise the substitute), wrapped greedily like a text box does.
"""
from __future__ import annotations

import copy
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT_DIR = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
FONT_FILES = {  # family -> (regular, bold) font files a viewer is likely to have
    "Arial": ("arial.ttf", "arialbd.ttf"), "Calibri": ("calibri.ttf", "calibrib.ttf"),
    "Georgia": ("georgia.ttf", "georgiab.ttf"), "Verdana": ("verdana.ttf", "verdanab.ttf"),
    "Segoe UI": ("segoeui.ttf", "segoeuib.ttf"), "Times New Roman": ("times.ttf", "timesbd.ttf"),
}
def theme_font() -> str:
    """PowerPoint renders an unavailable font with the deck theme's body font. Read it from the template we export with."""
    prs = Presentation()
    part = next(r.target_part for r in prs.slide_masters[0].part.rels.values() if r.reltype.endswith("/theme"))
    m = re.search(rb'<a:minorFont>\s*<a:latin typeface="([^"]+)"', part.blob)
    return m.group(1).decode() if m else "Calibri"


SUBSTITUTE = theme_font()     # "Calibri" for the default template; confirmed by rendering in PowerPoint
MARGIN_X, MARGIN_Y = 0.1, 0.05  # PowerPoint default text box insets, inches
LINE = 1.2                       # single-spacing line height ≈ 1.2 × size
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL, "roundrect": MSO_SHAPE.ROUNDED_RECTANGLE}


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


# ------------------------------------------------------------------ fonts + measurement

def installed(family: str) -> bool:
    return family in FONT_FILES and (FONT_DIR / FONT_FILES[family][0]).exists()


def rendered_family(family: str) -> str:
    return family if installed(family) else SUBSTITUTE


def font(family: str, size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    reg, bd = FONT_FILES[rendered_family(family)]
    return ImageFont.truetype(str(FONT_DIR / (bd if bold else reg)), max(1, round(size_pt * 96 / 72)))


def wrap(text: str, f: ImageFont.FreeTypeFont, width_px: float) -> list[str]:
    lines, cur = [], ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if f.getlength(trial) <= width_px or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    return lines + ([cur] if cur else [])


def fits(layer: dict, size: float | None = None) -> tuple[bool, int, float]:
    """(fits, line count, needed height in) for a text layer at `size` pt."""
    size = size or layer["size"]
    x, y, w, h = layer["box"]
    f = font(layer["font"], size, layer.get("bold", False))
    lines = wrap(layer["text"], f, (w - 2 * MARGIN_X) * 96)
    widest = max(f.getlength(l) for l in lines) / 96
    need = len(lines) * size * LINE / 72 + 2 * MARGIN_Y
    return need <= h + 1e-6 and widest <= w - 2 * MARGIN_X + 1e-6, len(lines), need


# ------------------------------------------------------------------ colour

def _lab(hexc: str):
    r, g, b = (int(hexc[i:i + 2], 16) / 255 for i in (1, 3, 5))
    lin = lambda c: c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = lin(r), lin(g), lin(b)
    X, Y, Z = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047, 0.2126 * r + 0.7152 * g + 0.0722 * b, (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    return 116 * f(Y) - 16, 500 * (f(X) - f(Y)), 200 * (f(Y) - f(Z))


def delta_e(a: str, b: str) -> float:
    (l1, a1, b1), (l2, a2, b2) = _lab(a), _lab(b)
    return ((l1 - l2) ** 2 + (a1 - a2) ** 2 + (b1 - b2) ** 2) ** 0.5


def nearest_brand(c: str, palette: dict) -> tuple[str, str, float]:
    name, hexc = min(palette.items(), key=lambda kv: delta_e(c, kv[1]))
    return name, hexc, delta_e(c, hexc)


# ------------------------------------------------------------------ lint

@dataclass
class Issue:
    slide: str
    layer: str
    code: str
    detail: str
    fix: str = ""


def lint(scene: dict) -> list[Issue]:
    brand, out = scene["brand"], []
    for s in scene["slides"]:
        for L in s["layers"]:
            sid, lid = s["id"], L["id"]
            for key in ("fill", "color"):
                if key in L and L[key].upper() not in {v.upper() for v in brand["colors"].values()}:
                    n, hx, de = nearest_brand(L[key], brand["colors"])
                    if de > 10:
                        out.append(Issue(sid, lid, "OFF_BRAND_COLOR", f"{L[key]} is ΔE {de:.0f} from the nearest brand colour ({n} {hx})", f"snap to {n} {hx}"))
            if L["type"] == "shape":
                if L.get("blend") and L["blend"] != "normal":
                    out.append(Issue(sid, lid, "UNSUPPORTED_EFFECT", f"blend mode '{L['blend']}' has no PPTX equivalent — exports as a plain opaque fill",
                                     "replace with solid fill + transparency"))
                if L.get("backdrop_blur"):
                    out.append(Issue(sid, lid, "UNSUPPORTED_EFFECT", f"backdrop blur {L['backdrop_blur']}px has no PPTX equivalent — the card becomes a flat panel",
                                     "replace with solid fill + transparency"))
                continue
            if L["type"] != "text":
                continue
            if L["font"] not in brand["fonts"] and L["font"] not in FONT_FILES:
                out.append(Issue(sid, lid, "OFF_BRAND_FONT", f"'{L['font']}' isn't a brand font"))
            if not installed(L["font"]):
                out.append(Issue(sid, lid, "FONT_SUBSTITUTED", f"'{L['font']}' isn't installed/embeddable on a typical viewer machine — PowerPoint renders it in {SUBSTITUTE}",
                                 f"set font to brand fallback '{brand['fonts'].get(L['font'], SUBSTITUTE)}' explicitly"))
            if L["size"] < brand["min_text_pt"]:
                out.append(Issue(sid, lid, "TEXT_TOO_SMALL", f"{L['size']} pt is below the {brand['min_text_pt']} pt minimum for customer-facing decks",
                                 f"raise to {brand['min_text_pt']} pt"))
            ok, n, need = fits(L)
            if not ok:
                out.append(Issue(sid, lid, "TEXT_OVERFLOW", f"in {rendered_family(L['font'])} {L['size']} pt the text needs {n} lines / {need:.2f} in, the box is {L['box'][3]:.2f} in tall",
                                 "shrink to fit or enlarge the box"))
    return group_fonts(out)


def group_fonts(issues: list[Issue]) -> list[Issue]:
    """One FONT_SUBSTITUTED finding per font family (listing the layers), instead of one per layer."""
    other = [i for i in issues if i.code != "FONT_SUBSTITUTED"]
    fams: dict[str, list[Issue]] = {}
    for i in issues:
        if i.code == "FONT_SUBSTITUTED":
            fams.setdefault(i.detail.split("'")[1], []).append(i)
    grouped = [Issue(",".join(sorted({x.slide for x in v})), f"{len(v)} layers", "FONT_SUBSTITUTED",
                     v[0].detail + f" ({len(v)} text layers: " + ", ".join(f"{x.slide}/{x.layer}" for x in v) + ")", v[0].fix)
               for v in fams.values()]
    return grouped + other


# ------------------------------------------------------------------ fix

def fix(scene: dict) -> tuple[dict, list[str]]:
    s2, log = copy.deepcopy(scene), []
    brand = s2["brand"]
    for s in s2["slides"]:
        for L in s["layers"]:
            for key in ("fill", "color"):
                if key in L:
                    n, hx, de = nearest_brand(L[key], brand["colors"])
                    if L[key].upper() not in {v.upper() for v in brand["colors"].values()} and de > 10:
                        log.append(f"{s['id']}/{L['id']}: {key} {L[key]} → {n} {hx}")
                        L[key] = hx
            if L["type"] == "shape" and (L.get("blend") not in (None, "normal") or L.get("backdrop_blur")):
                L.pop("blend", None)
                L.pop("backdrop_blur", None)
                L["opacity"] = min(L.get("opacity", 1.0), 0.75)
                log.append(f"{s['id']}/{L['id']}: unsupported effect → solid fill at {L['opacity']:.0%} opacity")
            if L["type"] != "text":
                continue
            if not installed(L["font"]):
                fb = brand["fonts"].get(L["font"], SUBSTITUTE)
                log.append(f"{s['id']}/{L['id']}: font {L['font']} → {fb} (explicit fallback)")
                L["font"] = fb
            if L["size"] < brand["min_text_pt"]:
                log.append(f"{s['id']}/{L['id']}: {L['size']} pt → {brand['min_text_pt']} pt")
                L["size"] = brand["min_text_pt"]
            if not fits(L)[0]:
                size = L["size"]
                while size > brand["min_text_pt"] and not fits(L, size)[0]:
                    size -= 1
                if fits(L, size)[0]:
                    log.append(f"{s['id']}/{L['id']}: shrink {L['size']} → {size} pt to fit")
                    L["size"] = size
                else:
                    log.append(f"{s['id']}/{L['id']}: still overflows at {brand['min_text_pt']} pt — needs a human edit")
    return s2, log


# ------------------------------------------------------------------ export + round trip

def _alpha(fill_el, opacity: float):
    clr = fill_el._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
    clr.append(a)


def export(scene: dict, path: Path) -> None:
    prs = Presentation()
    w, h = scene["size_in"]
    prs.slide_width, prs.slide_height = Inches(w), Inches(h)
    for s in scene["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for L in s["layers"]:
            x, y, bw, bh = (Inches(v) for v in L["box"])
            if L["type"] == "shape":
                sh = slide.shapes.add_shape(SHAPES[L["shape"]], x, y, bw, bh)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor.from_string(L["fill"][1:])
                if L.get("opacity", 1) < 1:
                    _alpha(sh.fill, L["opacity"])
                sh.line.fill.background()
                sh.shadow.inherit = False
            else:
                sh = slide.shapes.add_textbox(x, y, bw, bh)
                tf = sh.text_frame
                tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = Inches(MARGIN_X)
                tf.margin_top = tf.margin_bottom = Inches(MARGIN_Y)
                run = tf.paragraphs[0].add_run()
                run.text = L["text"]
                run.font.name, run.font.size, run.font.bold = L["font"], Pt(L["size"]), L.get("bold", False)
                run.font.color.rgb = RGBColor.from_string(L["color"][1:])
            sh.name = f"{s['id']}/{L['id']}"
    prs.save(path)


def round_trip(scene: dict, path: Path) -> list[Issue]:
    """Re-read the .pptx and check every layer survived with the same geometry, font and size."""
    prs, out = Presentation(path), []
    want = {f"{s['id']}/{L['id']}": (s["id"], L) for s in scene["slides"] for L in s["layers"]}
    seen = set()
    for slide in prs.slides:
        for sh in slide.shapes:
            sid, L = want[sh.name]
            seen.add(sh.name)
            exp = [int(Inches(v)) for v in L["box"]]
            got = [sh.left, sh.top, sh.width, sh.height]
            if max(abs(a - b) for a, b in zip(exp, got)) > 1:
                out.append(Issue(sid, L["id"], "GEOMETRY_DRIFT", f"expected {exp} EMU, got {got}"))
            if L["type"] == "text":
                r = sh.text_frame.paragraphs[0].runs[0]
                if r.font.name != L["font"] or abs(r.font.size.pt - L["size"]) > 0.01:
                    out.append(Issue(sid, L["id"], "TEXT_PROPS_CHANGED", f"{r.font.name} {r.font.size.pt}pt vs {L['font']} {L['size']}pt"))
    for k in set(want) - seen:
        out.append(Issue(want[k][0], want[k][1]["id"], "LAYER_LOST", "layer missing from the .pptx"))
    return out
